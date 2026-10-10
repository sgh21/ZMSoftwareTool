"""窝深理论值与误差报警设置；按当前设置评价数据，不改原始测量。"""

import csv
from datetime import datetime
from math import isclose, isfinite
from pathlib import Path
from statistics import fmean, pstdev, pvariance

from core.services.position_persistence import read_json, write_document


DEFAULT_SETTINGS = {
    "depth_mode": "per_hole",
    "uniform_depth": 1.5,
    "error_lower_mm": None,
    "error_upper_mm": None,
}


def validate_feed_depth_settings(settings):
    values = {key: settings.get(key, default) for key, default in DEFAULT_SETTINGS.items()}
    if values["depth_mode"] not in ("per_hole", "uniform"):
        raise ValueError("请选择逐孔理论值或统一理论窝深。")
    if values["depth_mode"] == "uniform" or values["uniform_depth"] is not None:
        try:
            depth = float(values["uniform_depth"])
        except (TypeError, ValueError):
            raise ValueError("理论窝深必须为大于 0 的有限数值，单位 mm。") from None
        if not isfinite(depth) or depth <= 0:
            raise ValueError("理论窝深必须为大于 0 的有限数值，单位 mm。")
        values["uniform_depth"] = depth
    lower, upper = values["error_lower_mm"], values["error_upper_mm"]
    if lower is None and upper is None:
        return values
    try:
        lower, upper = float(lower), float(upper)
    except (TypeError, ValueError):
        raise ValueError("请同时填写误差下限和上限，或同时留空。") from None
    if not isfinite(lower) or not isfinite(upper):
        raise ValueError("误差下限和上限必须为有限数值，单位 mm。")
    if lower > upper:
        raise ValueError("误差下限不能大于上限。")
    values.update(error_lower_mm=lower, error_upper_mm=upper)
    return values


def load_feed_depth_settings(path):
    path = Path(path)
    return validate_feed_depth_settings(read_json(path)) if path.exists() else dict(DEFAULT_SETTINGS)


def save_feed_depth_settings(path, settings):
    values = validate_feed_depth_settings(settings)
    write_document(path, values)
    return values


def load_feed_depth_data(path, settings, progress=None):
    """读取一文件一批的 CSV / Excel；缺测留空，保留原文件和模拟来源标记。"""
    path = Path(path)
    settings = validate_feed_depth_settings(settings)
    if progress:
        progress(0, f"正在读取 {path.name}")
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            fields = reader.fieldnames or []
            records = list(reader)
    elif path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            source = workbook.worksheets[0].iter_rows(values_only=True)
            fields = list(next(source, ()))
            records = [dict(zip(fields, row)) for row in source]
        finally:
            workbook.close()
    else:
        raise ValueError("请选择 CSV 或 Excel (.xlsx) 窝深数据。")
    required = {"hole_id", "actual_depth_mm"}
    if settings["depth_mode"] == "per_hole":
        required.add("theoretical_depth_mm")
    missing = required - set(fields)
    if missing:
        raise ValueError(f"缺少字段：{', '.join(sorted(missing))}；请查看“数据格式”。")

    def number(value, field, line):
        if value is None or str(value).strip() == "":
            return None
        try:
            result = float(value)
        except (TypeError, ValueError):
            raise ValueError(f"第 {line} 行 {field} 必须为有限数值。") from None
        if not isfinite(result):
            raise ValueError(f"第 {line} 行 {field} 必须为有限数值。")
        return result

    rows, identities = [], set()
    for line, record in enumerate(records, 2):
        if all(value is None or str(value).strip() == "" for value in record.values()):
            continue
        text = {key: str(value).strip() if value is not None else "" for key, value in record.items()}
        hole_id, row_id = text.get("hole_id", ""), text.get("row_id", "")
        if not hole_id:
            raise ValueError(f"第 {line} 行缺少 hole_id。")
        if (row_id, hole_id) in identities:
            raise ValueError(f"第 {line} 行孔号重复：{row_id} / {hole_id}。")
        identities.add((row_id, hole_id))
        actual = number(record.get("actual_depth_mm"), "actual_depth_mm", line)
        try:
            theory = number(record.get("theoretical_depth_mm"), "theoretical_depth_mm", line)
        except ValueError:
            if settings["depth_mode"] == "per_hole":
                raise
            theory = None
        if settings["depth_mode"] == "per_hole" and (theory is None or theory <= 0):
            raise ValueError(f"第 {line} 行 theoretical_depth_mm 必须大于 0。")
        measured_at = text.get("measured_at", "")
        if measured_at:
            try:
                measured_at = datetime.fromisoformat(measured_at).isoformat()
            except ValueError:
                raise ValueError(f"第 {line} 行 measured_at 请使用 ISO 日期或时间，例如 2026-10-10T09:00:00。") from None
        simulated = text.get("data_source") == "simulation" or text.get("is_simulated", "").lower() in ("true", "1")
        rows.append({
            "hole_id": hole_id, "row_id": row_id,
            "actual_depth_mm": actual, "theoretical_depth_mm": theory,
            "batch_id": text.get("batch_id") or path.stem,
            "measured_at": measured_at, "condition_id": text.get("condition_id", ""),
            "data_source": "simulation" if simulated else "measurement", "is_simulated": simulated,
        })
    if not rows:
        raise ValueError("文件中没有孔位数据。")
    for field in ("batch_id", "condition_id", "is_simulated"):
        if len({row[field] for row in rows}) != 1:
            raise ValueError(f"每个文件只能包含一批数据，{field} 必须整列一致。")
    if not any(row["actual_depth_mm"] is not None for row in rows):
        raise ValueError("文件中没有有效实测窝深；缺测不会补零，请检查 actual_depth_mm。")
    if progress:
        progress(70, f"已读取 {len(rows)} 孔，正在准备分组")
    return {"rows": rows, "history": [{key: rows[0][key] for key in
            ("batch_id", "condition_id", "data_source", "is_simulated")}], "source_path": str(path)}


def _over_limit(bias, settings):
    lower, upper = settings["error_lower_mm"], settings["error_upper_mm"]
    if lower is None:
        return False
    return ((bias < lower and not isclose(bias, lower, rel_tol=0.0, abs_tol=1e-12))
            or (bias > upper and not isclose(bias, upper, rel_tol=0.0, abs_tol=1e-12)))


def evaluate_feed_depth(data, settings):
    """按批次、孔排和理论值分组；缺测不补零，不改写原始测量。"""
    values = validate_feed_depth_settings(settings)
    rows, history = [], []
    batches = {}
    metadata = {record["batch_id"]: record for record in data["history"]}
    for source in data["rows"]:
        theory = values["uniform_depth"] if values["depth_mode"] == "uniform" else source["theoretical_depth_mm"]
        if theory is None or theory <= 0:
            raise ValueError("数据缺少有效逐孔理论窝深，请补充理论值或选择统一理论窝深。")
        actual = source["actual_depth_mm"]
        bias = actual - theory if actual is not None else None
        row = {**source, "theoretical_depth_mm": theory, "bias_mm": bias,
               "alarm": bias is not None and _over_limit(bias, values)}
        rows.append(row)
        key = (row["batch_id"], row.get("row_id", ""), theory)
        batches.setdefault(key, []).append(row)
    for (batch_id, row_id, theory), batch in batches.items():
        source = metadata[batch_id]
        depths = [row["actual_depth_mm"] for row in batch if row["actual_depth_mm"] is not None]
        bias = fmean(row["bias_mm"] for row in batch if row["bias_mm"] is not None) if depths else None
        times = [row["measured_at"] for row in batch if row.get("measured_at")]
        history.append({
            **source,
            "row_id": row_id, "theoretical_depth_mm": theory,
            "measured_start_at": min(times, key=lambda value: datetime.fromisoformat(value).timestamp()) if times else source.get("measured_start_at", ""),
            "measured_at": max(times, key=lambda value: datetime.fromisoformat(value).timestamp()) if times else source.get("measured_at", ""),
            "count": len(depths), "mean_depth_mm": fmean(depths) if depths else None,
            "variance_mm2": pvariance(depths) if depths else None,
            "stddev_mm": pstdev(depths) if depths else None,
            "bias_mm": bias,
            "alarm": bias is not None and _over_limit(bias, values),
            "alarm_count": sum(row["alarm"] for row in batch),
        })
    return {"rows": rows, "history": history}
